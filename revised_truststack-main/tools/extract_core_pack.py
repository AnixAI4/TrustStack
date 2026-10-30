#!/usr/bin/env python3
"""Extract the core pack, the 16 domain weights and the 98 sample responses.

The workbook is the scoring oracle (D-007) and the only source for the core
pack, so this runs once whenever the workbook changes and its OUTPUT is
committed. The seeder reads that JSON; nothing at runtime opens an .xlsx, and
`packages/core` never gains an xlsx dependency.

Three outputs:
  * `core-pack.json` — the 98 authored core questions + the `Config` weights.
  * `workbook-sample.json` — the `Responses` sheet's 98 answered rows, which
    are the parity gate's fixture (Slice 5, D-048). The engine runs them TWICE,
    under the workbook's naive `MIN(claimed, evidence)` and under the spec's
    cap table, and must reproduce both sets of numbers exactly.
  * `ropa-register.json` — the `ROPA_Register` sheet's ten records, exactly
    as typed (Slice 7). The seeder loads them for Mumbai University. They are
    NOT an oracle: the sheet computes nothing, and its Risk Rating is entered
    by hand.

Two things this script refuses to do:
  * read `Question_Bank.Domain Weight` — it is stale and totals 104. Weights
    come from `Config` only (D-006), and the total must be exactly 100.
  * keep the generic `Legal / Control Basis` sentence — it is a placeholder,
    not a citation, so it is stored as null (D-013).

Usage:  python3 tools/extract_core_pack.py [--check]
        --check verifies the committed JSON matches the workbook; exit 1 if not.
"""
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORKBOOK = os.path.join(ROOT, "docs", "source",
                        "TrustStack_DPDP_Compass_Excel_MVP_Validated_Step12.xlsx")
OUT = os.path.join(ROOT, "packages", "db", "question-bank", "core-pack.json")
SAMPLE_OUT = os.path.join(ROOT, "packages", "db", "question-bank", "workbook-sample.json")
ROPA_OUT = os.path.join(ROOT, "packages", "db", "question-bank", "ropa-register.json")

DOMAIN_CODES = ["GOV", "INV", "ROPA", "PURPOSE", "NOTICE", "CONSENT", "RIGHTS", "GRIEVANCE",
                "RETENTION", "ACCURACY", "SECURITY", "BREACH", "VENDOR", "CHILD",
                "CROSSBORDER", "SDF_AI"]
DOMAIN_BY_ID = {f"D{i:02d}": c for i, c in enumerate(DOMAIN_CODES, start=1)}

LEGAL_BASIS_PLACEHOLDER = ("Mapped to DPDP Act 2023 / DPDP Rules 2025 "
                           "operational readiness requirement.")
# Documented equivalence, not new content: docs/QUESTION_FORMAT.md maps
# P0->critical, P1->high, P2->medium. The core pack authored Criticality
# instead of Priority, so the same mapping is applied in reverse.
PRIORITY_BY_CRITICALITY = {"Critical": "P0", "High": "P1", "Medium": "P2", "Low": "P3"}

MANDATORY_NA_RULE = ("Mandatory for every organisation. Not Applicable requires a "
                     "written reason and reviewer approval.")
DEFAULT_NA_RULE = ("May be marked Not Applicable with a written reason; excluded from "
                   "the scoring denominator.")


def rows(ws):
    header = [c.value for c in ws[1]]
    header = [h for h in header if h]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        yield dict(zip(header, [ws.cell(r, c).value for c in range(1, len(header) + 1)]))


def text(value):
    s = str(value).strip() if value is not None else ""
    return s or None


def parse_applicability(raw):
    """`Default Applicability` -> (applicability, trigger_text).

    "Mandatory where vendors exist" is conditional with a trigger: a question
    that applies only when vendors exist is not mandatory for everyone, and
    D13's own default applicability in `Config` agrees.
    """
    cell = (raw or "").strip()
    low = cell.lower()
    if low.startswith("mandatory where ") or low.startswith("mandatory if "):
        return "conditional", cell.split(" ", 2)[2].strip()
    if low == "conditional":
        return "conditional", None
    return "mandatory", None


def build():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    config = wb["Config"]
    domains = []
    for r in range(33, 49):
        domain_id = config.cell(r, 1).value
        if not domain_id:
            continue
        domains.append({
            "domainId": domain_id,
            "code": DOMAIN_BY_ID[domain_id],
            "name": str(config.cell(r, 2).value).strip(),
            "weight": float(config.cell(r, 3).value),
            "sortOrder": len(domains) + 1,
        })
    if len(domains) != 16:
        sys.exit(f"::error::expected 16 Config domain rows, found {len(domains)}")
    total = sum(d["weight"] for d in domains)
    if round(total, 2) != 100:
        sys.exit(f"::error::DOMAIN_WEIGHT_TOTAL_INVALID: Config weights total {total}")

    questions = []
    for row in rows(wb["Question_Bank"]):
        applicability, trigger_text = parse_applicability(row["Default Applicability"])
        criticality = text(row["Criticality"])
        questions.append({
            "questionCode": str(row["Question ID"]).strip(),
            "questionText": str(row["Question"]).strip(),
            "domainCode": DOMAIN_BY_ID[str(row["Domain ID"]).strip()],
            "secondaryDomainCodes": [],
            "applicability": applicability,
            "typeText": text(row["Default Applicability"]),
            "triggerText": trigger_text,
            "riskTags": [],
            "evidenceGuidance": text(row["Evidence Required"]),
            "riskIfWeak": None,
            "suggestedOwner": text(row["Default Owner"]),
            "priority": PRIORITY_BY_CRITICALITY.get(criticality or ""),
            "criticality": criticality,
            "legalControlBasis": (None if (row["Legal / Control Basis"] or "").strip()
                                  == LEGAL_BASIS_PLACEHOLDER
                                  else text(row["Legal / Control Basis"])),
            "recommendedFix": text(row["Recommended Fix"]),
            "truststackModule": text(row["TrustStack Module"]),
            "naRule": MANDATORY_NA_RULE if applicability == "mandatory" else DEFAULT_NA_RULE,
        })

    codes = [q["questionCode"] for q in questions]
    if len(set(codes)) != len(codes):
        sys.exit("::error::duplicate question code in Question_Bank")
    missing = [d["code"] for d in domains
               if d["code"] not in {q["domainCode"] for q in questions}]
    if missing:
        sys.exit(f"::error::core pack covers no question for: {', '.join(missing)}")

    return {
        "_comment": ("Generated by tools/extract_core_pack.py from "
                     "docs/source/TrustStack_DPDP_Compass_Excel_MVP_Validated_Step12.xlsx. "
                     "Do not edit by hand. Weights come from Config (D-006); the "
                     "Question_Bank weight column is stale at 104 and is never read."),
        "packCode": "CORE_DPDP_V1",
        "sourceTag": "core",
        "version": "1.0",
        "status": "draft",
        "domains": domains,
        "questions": questions,
    }


def build_sample(domains):
    """The `Responses` sheet as the parity fixture (D-048).

    Six columns only — everything the engine needs and nothing it must not see.
    The workbook's own `Adjusted Score`, `Gap Severity` and `Evidence
    Confidence %` columns are deliberately NOT carried: the gate exists to make
    our engine derive them, not to copy them.
    """
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb["Responses"]

    samples = []
    for row in rows(ws):
        domain_id = str(row["Domain ID"]).strip()
        if domain_id not in DOMAIN_BY_ID:
            sys.exit(f"::error::Responses row {row['Question ID']} has unknown domain {domain_id}")
        samples.append({
            "questionCode": str(row["Question ID"]).strip(),
            "domainId": domain_id,
            "domainCode": DOMAIN_BY_ID[domain_id],
            "applicability": str(row["Applicability"]).strip(),
            "claimedMaturity": int(row["Maturity Score (0-5)"]),
            "evidenceLevel": int(row["Evidence Level (0-5)"]),
            "evidenceCurrent": str(row["Evidence Current"]).strip(),
        })

    codes = [s["questionCode"] for s in samples]
    if len(set(codes)) != len(codes):
        sys.exit("::error::duplicate question code in Responses")
    for s in samples:
        for field in ("claimedMaturity", "evidenceLevel"):
            if not 0 <= s[field] <= 5:
                sys.exit(f"::error::{s['questionCode']}.{field} is {s[field]}, expected 0–5")
    covered = {s["domainCode"] for s in samples}
    missing = [d["code"] for d in domains if d["code"] not in covered]
    if missing:
        # Parity depends on all 16 domains being answered: the overall evidence
        # confidence is a mean over SCORED domains (D-048, §4.4.3).
        sys.exit(f"::error::the sample answers no question for: {', '.join(missing)}")

    return {
        "_comment": ("Generated by tools/extract_core_pack.py from the Responses sheet of "
                     "docs/source/TrustStack_DPDP_Compass_Excel_MVP_Validated_Step12.xlsx. "
                     "Do not edit by hand. This is the scoring engine's parity fixture "
                     "(D-007, D-048): the engine runs it under the workbook's naive "
                     "MIN(claimed, evidence) and under the spec's cap table, and must "
                     "reproduce `expected` exactly. The workbook's own Adjusted Score and "
                     "Gap Severity columns are NOT carried — the engine derives them."),
        "source": "Responses",
        "assessment": "Mumbai University DPDP Compass Baseline",
        "expected": EXPECTED,
        "responses": samples,
    }


# The two runs the CI gate asserts (D-048), read from the workbook's cached
# Dashboard values on 2026-09-20 and re-derived here. `naive` proves the
# AGGREGATION matches the workbook; `spec` pins the product's own numbers. The
# only difference is the cap, on the five CB1.x rows (claim 2 + E1).
EXPECTED = {
    "naive": {"rawScore": 24.7014, "evidenceConfidence": 23.9, "criticalCount": 10, "highCount": 57},
    "spec": {"rawScore": 25.3014, "evidenceConfidence": 23.9, "criticalCount": 10, "highCount": 52},
    # Red flags cap the sample at 50 (CHILD) under both rules and both
    # thresholds — so the raw score, far below it, is also the final score.
    "scoreCap": 50,
}


# The 24 columns of `ROPA_Register`, in sheet order. A rename in the workbook
# fails the extraction rather than silently dropping a column.
ROPA_COLUMNS = [
    "ROPA ID", "Processing Activity", "Department / Owner", "Data Principal Type",
    "Personal Data Categories", "Purpose", "DPDP Basis / Legal Review", "Collection Source",
    "System(s)", "Processor / Vendor", "Sharing / Disclosure", "Retention Rule",
    "Security Safeguards", "Notice Linked", "Consent Linked", "Rights Workflow",
    "Grievance Linked", "Children Data?", "Cross-Border?", "AI / Automated?", "Evidence",
    "Risk Rating", "Last Reviewed", "Next Review",
]


def cell_text(value):
    """A cell as the text a CSV export of it would carry. Dates are calendar dates."""
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def build_ropa():
    """The ten processing records, cell for cell (Slice 7).

    `data_only=False`: a formula would come through as its formula text and
    fail the check below — the sheet has none, and the importer must never
    evaluate one either.
    """
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    ws = wb["ROPA_Register"]
    header = [cell_text(c.value) for c in ws[1]][:len(ROPA_COLUMNS)]
    if header != ROPA_COLUMNS:
        sys.exit(f"::error::ROPA_Register columns changed: {header}")

    records = []
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, len(ROPA_COLUMNS) + 1)]
        if values[0] is None:
            continue
        for v in values:
            if isinstance(v, str) and v.startswith("="):
                sys.exit(f"::error::ROPA_Register row {r} holds a formula — the register is typed text")
        records.append(dict(zip(ROPA_COLUMNS, [cell_text(v) for v in values])))

    codes = [rec["ROPA ID"] for rec in records]
    if len(set(codes)) != len(codes):
        sys.exit("::error::duplicate ROPA ID in ROPA_Register")
    return {
        "_comment": ("Generated by tools/extract_core_pack.py from the ROPA_Register sheet of "
                     "docs/source/TrustStack_DPDP_Compass_Excel_MVP_Validated_Step12.xlsx. "
                     "Do not edit by hand. The ten records exactly as typed, loaded for "
                     "Mumbai University by the seeder through the same validation an import "
                     "uses. NOT an oracle: the sheet computes no completeness score and its "
                     "Risk Rating is entered by hand, so nothing may claim to match it."),
        "source": "ROPA_Register",
        "columns": ROPA_COLUMNS,
        "records": records,
    }


def write_or_check(path, payload, label):
    text_out = json.dumps(payload, indent=2, ensure_ascii=False) + "\n"
    if "--check" in sys.argv:
        if not os.path.exists(path):
            print(f"::error::{path} is missing — run: python3 tools/extract_core_pack.py")
            return 1
        with open(path, encoding="utf-8") as fh:
            if fh.read() != text_out:
                print(f"::error::{label} is stale — run: python3 tools/extract_core_pack.py")
                return 1
        return 0
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text_out)
    return 0


def main():
    bank = build()
    sample = build_sample(bank["domains"])

    ropa = build_ropa()

    failed = write_or_check(OUT, bank, "core-pack.json")
    failed |= write_or_check(SAMPLE_OUT, sample, "workbook-sample.json")
    failed |= write_or_check(ROPA_OUT, ropa, "ropa-register.json")
    if failed:
        return 1

    if "--check" in sys.argv:
        print(f"core-pack.json is current ({len(bank['questions'])} questions, "
              f"{len(bank['domains'])} domains)")
        print(f"workbook-sample.json is current ({len(sample['responses'])} sample responses)")
        print(f"ropa-register.json is current ({len(ropa['records'])} processing records)")
        return 0
    print(f"wrote {OUT}: {len(bank['questions'])} questions, {len(bank['domains'])} domains, "
          f"weights total {sum(d['weight'] for d in bank['domains']):g}")
    print(f"wrote {SAMPLE_OUT}: {len(sample['responses'])} sample responses")
    print(f"wrote {ROPA_OUT}: {len(ropa['records'])} processing records")
    return 0


if __name__ == "__main__":
    sys.exit(main())
